""" Workspace directory indexing for retrieval (RAG).

The indexer walks a workspace's attached directories and ingests indexable
files into a per-directory sidecar store
(``<dir>/.subconscious/<workspace_uuid>/index.db``). For each file it extracts
text, splits it into overlapping chunks, computes an embedding per chunk, and
extracts a structural knowledge graph (imports / definitions).

Everything here is **synchronous** and designed to run inside a worker thread
(see :meth:`Engine.reindex_workspace`) so the asyncio event loop — and therefore
the desktop UI's rendering — is never blocked by file I/O, PDF/Office parsing,
hashing, or SQLite writes.
"""
from __future__ import annotations

import os
import hashlib
import logging
import pathlib
import threading
from typing import Callable, Optional

from .rag.sidecar import SidecarStore, SIDECAR_DIRNAME
from .rag.embeddings import Embedder, pack_vector, get_default_embedder
from .rag import graph as kg


logger = logging.getLogger("subconscious")


# File types we can extract text from.
_PLAIN_EXTS = {
  ".txt", ".md", ".markdown", ".rst", ".py", ".js", ".ts", ".tsx", ".jsx",
  ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".csv", ".html", ".htm",
  ".css", ".scss", ".java", ".c", ".cc", ".cpp", ".h", ".hpp", ".go", ".rs",
  ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".xml", ".log", ".tex",
}
_STRUCTURED_EXTS = {".pdf", ".docx", ".xlsx"}

# Directories we never descend into (our own sidecars, VCS/build noise).
_SKIP_DIRS = {
  SIDECAR_DIRNAME, ".git", ".hg", ".svn", "__pycache__", "node_modules",
  ".venv", "venv", ".mypy_cache", ".pytest_cache", ".hypothesis", ".idea", ".vscode",
}

_CHUNK_CHARS = 1500          # target characters per chunk
_CHUNK_OVERLAP = 200         # overlap between consecutive chunks
_MAX_FILE_BYTES = 20_000_000  # skip files larger than 20 MB
_HASH_LIMIT = 5_000_000       # only content-hash files up to 5 MB; larger rely on size+mtime


ProgressCb = Callable[[int, int, str], None]


class WorkspaceIndexer:
  """Ingests workspace directories into per-directory sidecar stores."""

  def __init__(self, embedder: Optional[Embedder] = None, cache_dir: Optional[str] = None):
    self.embedder = embedder or get_default_embedder()
    # Central fallback location for sidecars when an attached directory is not
    # writable (read-only mount / network share / restricted permissions).
    self.cache_dir = cache_dir

  # ------------------------------------------------------------------
  # Public (synchronous — run me in a worker thread)
  # ------------------------------------------------------------------

  def reindex_sync(
    self,
    workspace_uuid: str,
    directories: list[str],
    cancel: Optional[threading.Event] = None,
    progress: Optional[ProgressCb] = None,
  ) -> dict:
    """Incrementally (re)index every indexable file under *directories*.

    Runs synchronously; safe to call via ``loop.run_in_executor``. Honors the
    *cancel* event between files and reports via the *progress* callback.
    """
    cancel = cancel or threading.Event()

    # Collect the full work list first so progress totals are accurate.
    tasks: list[tuple[pathlib.Path, pathlib.Path]] = []
    for d in directories:
      root = pathlib.Path(d)
      if not root.exists() or not root.is_dir():
        continue
      for p in self._walk(root):
        tasks.append((root, p))

    total = len(tasks)
    if progress:
      progress(0, total, f"Found {total} files")

    stores: dict[str, SidecarStore] = {}
    seen: dict[str, set[str]] = {}
    indexed = 0

    try:
      for i, (root, p) in enumerate(tasks):
        if cancel.is_set():
          logger.info("Indexing cancelled after %d/%d files", i, total)
          break
        root_str = str(root)
        if root_str not in stores:
          # First file under this root: open (or fail) its sidecar once.
          try:
            stores[root_str] = SidecarStore.open(
              root, workspace_uuid, self.embedder.name, self.cache_dir
            )
            seen[root_str] = set()
          except OSError as exc:
            # No writable location (in-dir or cache) — skip this directory but
            # keep indexing the rest. None marks the root as unindexable.
            logger.warning("Skipping unindexable directory %s: %s", root, exc)
            stores[root_str] = None  # type: ignore[assignment]
        store = stores[root_str]
        if store is None:
          continue

        rel = os.path.relpath(str(p), root_str)
        seen[root_str].add(rel)
        try:
          if self._index_file(store, root, p, rel):
            indexed += 1
        except Exception as exc:  # never let one bad file kill the whole run
          logger.warning("Indexing error for %s: %s", p, exc)
          try:
            store.mark_error(rel, str(exc))
          except Exception:
            pass
        if progress:
          progress(i + 1, total, p.name)

      # Prune documents whose files were removed (only for fully-scanned roots).
      if not cancel.is_set():
        for root_str, store in stores.items():
          if store is not None:
            store.prune(seen.get(root_str, set()))
    finally:
      for store in stores.values():
        if store is not None:
          store.close()

    summary = {"indexed": indexed, "total": total, "cancelled": cancel.is_set()}
    return summary

  # ------------------------------------------------------------------
  # Walking & filtering
  # ------------------------------------------------------------------

  def _walk(self, root: pathlib.Path):
    """Yield indexable files under *root*, skipping noise directories."""
    for dirpath, dirnames, filenames in os.walk(root):
      # Prune skip dirs in-place so os.walk doesn't descend into them.
      dirnames[:] = [d for d in dirnames if d not in _SKIP_DIRS]
      for name in filenames:
        p = pathlib.Path(dirpath) / name
        if self._is_indexable(p):
          yield p

  @staticmethod
  def _is_indexable(p: pathlib.Path) -> bool:
    ext = p.suffix.lower()
    if ext not in _PLAIN_EXTS and ext not in _STRUCTURED_EXTS:
      return False
    try:
      return p.is_file() and p.stat().st_size <= _MAX_FILE_BYTES
    except OSError:
      return False

  # ------------------------------------------------------------------
  # Per-file indexing
  # ------------------------------------------------------------------

  def _index_file(self, store: SidecarStore, root: pathlib.Path, p: pathlib.Path, rel: str) -> bool:
    """Index a single file. Returns True when (re)chunked, False when skipped."""
    stat = p.stat()
    size = stat.st_size
    mtime = int(stat.st_mtime)

    existing = store.get_document(rel)
    content_hash = self._hash_file(p) if size <= _HASH_LIMIT else None

    # Unchanged since last index → skip.
    if existing and existing["status"] == "indexed" and existing["mtime"] == mtime and existing["size"] == size:
      if content_hash is None or existing["content_hash"] == content_hash:
        return False

    text = self._extract_text(p)
    chunks = self._chunk_text(text)

    # Embed each chunk (offline hashing embedder → cheap, thread-safe).
    vectors = [pack_vector(self.embedder.embed(c[0])) for c in chunks]

    doc_id, chunk_ids = store.replace_document(
      rel_path=rel,
      size=size,
      mtime=mtime,
      content_hash=content_hash,
      chunks=chunks,
      vectors=vectors,
      dim=self.embedder.dim,
    )

    # Structural knowledge graph: attach chunk-level provenance by mapping the
    # extractor's line numbers onto the chunk that covers that line.
    self._build_structural_kg(store, rel, text, p.suffix.lower(), doc_id, chunks, chunk_ids)
    return True

  def _build_structural_kg(
    self,
    store: SidecarStore,
    rel: str,
    text: str,
    ext: str,
    doc_id: int,
    chunks: list[tuple[str, Optional[int], Optional[int]]],
    chunk_ids: list[int],
  ) -> None:
    triples = kg.extract_structural(rel, text, ext)
    if not triples:
      return
    for t in triples:
      t["document_id"] = doc_id
      t["chunk_id"] = self._chunk_for_line(t.get("line"), chunks, chunk_ids)
      t.pop("line", None)
    try:
      store.add_triples(triples)
    except Exception as exc:
      logger.debug("KG add_triples failed for %s: %s", rel, exc)

  @staticmethod
  def _chunk_for_line(
    line: Optional[int],
    chunks: list[tuple[str, Optional[int], Optional[int]]],
    chunk_ids: list[int],
  ) -> Optional[int]:
    if line is None:
      return chunk_ids[0] if chunk_ids else None
    for idx, (_content, start, end) in enumerate(chunks):
      if start is not None and end is not None and start <= line <= end:
        return chunk_ids[idx]
    return chunk_ids[0] if chunk_ids else None

  # ------------------------------------------------------------------
  # Text extraction & chunking
  # ------------------------------------------------------------------

  @staticmethod
  def _hash_file(p: pathlib.Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
      for block in iter(lambda: f.read(65536), b""):
        h.update(block)
    return h.hexdigest()

  @staticmethod
  def _extract_text(p: pathlib.Path) -> str:
    ext = p.suffix.lower()
    if ext == ".docx":
      import docx as _docx
      doc = _docx.Document(str(p))
      return "\n".join(para.text for para in doc.paragraphs)
    if ext == ".xlsx":
      import openpyxl as _openpyxl
      wb = _openpyxl.load_workbook(p, data_only=True, read_only=True)
      rows = []
      for sheet in wb.worksheets:
        rows.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
          if any(cell is not None for cell in row):
            rows.append("\t".join(str(c) if c is not None else "" for c in row))
      return "\n".join(rows)
    if ext == ".pdf":
      import pypdf as _pypdf
      reader = _pypdf.PdfReader(str(p))
      pages = []
      for i, page in enumerate(reader.pages):
        t = page.extract_text()
        if t:
          pages.append(f"--- Page {i + 1} ---\n{t}")
      return "\n".join(pages)
    # Plain text (best-effort decode).
    return p.read_bytes().decode("utf-8", errors="replace")

  @staticmethod
  def _chunk_text(text: str) -> list[tuple[str, int, int]]:
    """Split text into overlapping character chunks with 1-based line ranges."""
    if not text.strip():
      return []

    # Build a sorted list of char offsets at the start of each line so a chunk's
    # char span can be mapped back to 1-based line numbers.
    offsets = [0]
    for i, ch in enumerate(text):
      if ch == "\n":
        offsets.append(i + 1)

    def line_of(offset: int) -> int:
      lo, hi = 0, len(offsets) - 1
      while lo < hi:
        mid = (lo + hi + 1) // 2
        if offsets[mid] <= offset:
          lo = mid
        else:
          hi = mid - 1
      return lo + 1  # 1-based

    chunks: list[tuple[str, int, int]] = []
    n = len(text)
    start = 0
    step = max(1, _CHUNK_CHARS - _CHUNK_OVERLAP)
    while start < n:
      end = min(n, start + _CHUNK_CHARS)
      content = text[start:end].strip()
      if content:
        chunks.append((content, line_of(start), line_of(max(start, end - 1))))
      start += step
    return chunks
