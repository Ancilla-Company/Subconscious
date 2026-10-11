"""Structured reading and conservative generation of common document files."""

from __future__ import annotations

import csv
from datetime import date, datetime
import os
from pathlib import Path
import textwrap
from typing import Any

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_OUTPUT_CHARS = 1_000_000
MAX_PAGES = 1_000
MAX_CELLS = 100_000


def read_document(path: Path) -> dict[str, Any]:
    if path.suffix.casefold() not in SUPPORTED:
        raise ValueError("Supported document types are PDF, DOCX, XLSX, and PPTX.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("The document exceeds the 50 MB limit.")
    readers = {".pdf": _read_pdf, ".docx": _read_docx,
               ".xlsx": _read_xlsx, ".pptx": _read_pptx}
    result = readers[path.suffix.casefold()](path)
    return {"schemaVersion": 1, "path": str(path), **result}


def _bounded(text: str, used: list[int]) -> str:
    remaining = MAX_OUTPUT_CHARS - used[0]
    if remaining <= 0:
        return ""
    value = text[:remaining]
    used[0] += len(value)
    return value


def _read_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDF files are unsupported.")
    if len(reader.pages) > MAX_PAGES:
        raise ValueError(f"The PDF exceeds the {MAX_PAGES}-page limit.")
    used = [0]
    pages = []
    warnings = []
    for number, page in enumerate(reader.pages, start=1):
        text = _bounded(page.extract_text() or "", used)
        if not text:
            warnings.append(f"Page {number} has no extractable text.")
        pages.append({"number": number, "text": text})
        if used[0] >= MAX_OUTPUT_CHARS:
            warnings.append("Extracted text was truncated at 1,000,000 characters.")
            break
    return {"format": "pdf", "content": {"pages": pages},
            "warnings": warnings, "truncated": used[0] >= MAX_OUTPUT_CHARS}


def _read_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    used = [0]
    paragraphs = [
        {"text": _bounded(item.text, used), "style": item.style.name if item.style else None}
        for item in document.paragraphs if item.text
    ]
    tables = []
    for table in document.tables:
        tables.append({"rows": [[_bounded(cell.text, used) for cell in row.cells]
                                for row in table.rows]})
    truncated = used[0] >= MAX_OUTPUT_CHARS
    return {"format": "docx", "content": {"paragraphs": paragraphs, "tables": tables},
            "warnings": ["Extracted text was truncated."] if truncated else [],
            "truncated": truncated}


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _read_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    cells = 0
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows():
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ValueError(f"The workbook exceeds the {MAX_CELLS}-cell limit.")
                rows.append([{"value": _cell_value(cell.value), "dataType": cell.data_type}
                             for cell in row])
            sheets.append({"name": worksheet.title, "rows": rows})
    finally:
        workbook.close()
    return {"format": "xlsx", "content": {"sheets": sheets},
            "warnings": ["Formulas are returned but are not calculated."], "truncated": False}


def _read_pptx(path: Path) -> dict[str, Any]:
    presentation = Presentation(path)
    if len(presentation.slides) > MAX_PAGES:
        raise ValueError(f"The presentation exceeds the {MAX_PAGES}-slide limit.")
    used = [0]
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts, tables = [], []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text:
                texts.append(_bounded(shape.text, used))
            if getattr(shape, "has_table", False):
                tables.append({"rows": [[_bounded(cell.text, used) for cell in row.cells]
                                        for row in shape.table.rows]})
        slides.append({"number": number, "texts": texts, "tables": tables})
        if used[0] >= MAX_OUTPUT_CHARS:
            break
    truncated = used[0] >= MAX_OUTPUT_CHARS
    return {"format": "pptx", "content": {"slides": slides},
            "warnings": ["Extracted text was truncated."] if truncated else [],
            "truncated": truncated}


def write_document(path: Path, content: Any, create_parents: bool = False) -> dict[str, Any]:
    extension = path.suffix.casefold()
    if extension not in SUPPORTED:
        raise ValueError("Supported document types are PDF, DOCX, XLSX, and PPTX.")
    if create_parents:
        path.parent.mkdir(parents=True, exist_ok=True)
    if not path.parent.is_dir():
        raise ValueError("The destination parent directory does not exist.")
    payload = content.get("content", content) if isinstance(content, dict) else content
    temporary = path.with_name(f".{path.stem}.subconscious-tmp{path.suffix}")
    writers = {".pdf": _write_pdf, ".docx": _write_docx,
               ".xlsx": _write_xlsx, ".pptx": _write_pptx}
    try:
        writers[extension](temporary, payload)
        read_document(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return {"status": "written", "path": str(path), "format": extension[1:]}


def _as_text(content: Any) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, dict):
        if "pages" in content:
            return "\n\n".join(str(page.get("text", "")) for page in content["pages"])
        if "paragraphs" in content:
            return "\n".join(str(item.get("text", item)) if isinstance(item, dict) else str(item)
                              for item in content["paragraphs"])
        if "slides" in content:
            return "\n\n---\n\n".join("\n".join(slide.get("texts", []))
                                          for slide in content["slides"])
    return str(content)


def _write_pdf(path: Path, content: Any) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 54
    for source_line in _as_text(content).splitlines() or [""]:
        for line in textwrap.wrap(source_line, 100, replace_whitespace=False) or [""]:
            if y < 54:
                pdf.showPage()
                y = height - 54
            pdf.drawString(54, y, line)
            y -= 14
    pdf.save()


def _write_docx(path: Path, content: Any) -> None:
    document = Document()
    if isinstance(content, dict):
        for item in content.get("paragraphs", []):
            text = item.get("text", "") if isinstance(item, dict) else str(item)
            paragraph = document.add_paragraph(text)
            style = item.get("style") if isinstance(item, dict) else None
            if style and style in document.styles:
                paragraph.style = style
        for table_data in content.get("tables", []):
            rows = table_data.get("rows", [])
            width = max((len(row) for row in rows), default=0)
            if rows and width:
                table = document.add_table(rows=len(rows), cols=width)
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        table.cell(row_index, column_index).text = str(value)
    else:
        for line in _as_text(content).splitlines():
            document.add_paragraph(line)
    document.save(path)


def _write_xlsx(path: Path, content: Any) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    if isinstance(content, str):
        sheets = [{"name": "Sheet1", "rows": list(csv.reader(content.splitlines()))}]
    elif isinstance(content, list):
        sheets = [{"name": "Sheet1", "rows": content}]
    else:
        sheets = content.get("sheets", [])
    if not sheets:
        sheets = [{"name": "Sheet1", "rows": []}]
    for sheet_index, sheet in enumerate(sheets, start=1):
        worksheet = workbook.create_sheet(str(sheet.get("name") or f"Sheet{sheet_index}")[:31])
        for row in sheet.get("rows", []):
            worksheet.append([cell.get("value") if isinstance(cell, dict) else cell for cell in row])
    workbook.save(path)


def _write_pptx(path: Path, content: Any) -> None:
    presentation = Presentation()
    if isinstance(content, dict):
        slides = content.get("slides", [])
    else:
        slides = []
        for block in _as_text(content).split("\n---\n"):
            lines = block.strip().splitlines()
            slides.append({"title": lines[0] if lines else "", "body": "\n".join(lines[1:])})
    if not slides:
        slides = [{"title": "", "body": ""}]
    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        texts = item.get("texts", [])
        title = item.get("title", texts[0] if texts else "")
        body = item.get("body", "\n".join(texts[1:]) if len(texts) > 1 else "")
        slide.shapes.title.text = str(title)
        slide.placeholders[1].text = str(body)
    presentation.save(path)
